"""Views Web Report Sales Srikaton."""
import datetime
from django.contrib.auth import login, logout, authenticate
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.db.models import Count
from django.http import HttpResponse, HttpResponseForbidden, JsonResponse
from django.shortcuts import render, redirect, get_object_or_404
from django.utils import timezone

from .models import User, Instansi, Kunjungan, FollowUpWA
from .forms import KunjunganForm
from . import wa_templates


# --- AUTH --------------------------------------------------------------------
def login_view(request):
    if request.user.is_authenticated:
        return redirect("dashboard")
    if request.method == "POST":
        u = authenticate(request,
                         username=request.POST.get("username"),
                         password=request.POST.get("password"))
        if u:
            login(request, u)
            return redirect("dashboard")
        messages.error(request, "Username atau password salah.")
    return render(request, "report/login.html")


def logout_view(request):
    logout(request)
    return redirect("login")


# --- DASHBOARD (cabang sesuai role) ------------------------------------------
@login_required
def dashboard(request):
    if request.user.is_admin_spv or request.user.is_superuser:
        return dashboard_spv(request)
    return dashboard_sales(request)


def dashboard_sales(request):
    hari_ini = timezone.localdate()
    kunjungan = Kunjungan.objects.filter(sales=request.user)
    ctx = {
        "total": kunjungan.count(),
        "bulan_ini": kunjungan.filter(
            tanggal__year=hari_ini.year, tanggal__month=hari_ini.month
        ).count(),
        "terbaru": kunjungan.select_related("instansi")[:5],
    }
    return render(request, "report/dashboard_sales.html", ctx)


def dashboard_spv(request):
    # Filter opsional
    bulan = request.GET.get("bulan", "")
    sales_id = request.GET.get("sales", "")
    status = request.GET.get("status", "")

    qs = Kunjungan.objects.select_related("sales", "instansi")
    if bulan:
        try:
            th, bl = bulan.split("-")
            qs = qs.filter(tanggal__year=int(th), tanggal__month=int(bl))
        except ValueError:
            pass
    if sales_id:
        qs = qs.filter(sales_id=sales_id)
    if status:
        qs = qs.filter(status=status)

    rekap_status = list(
        qs.values("status").annotate(jumlah=Count("id")).order_by("-jumlah")
    )
    rekap_sales = list(
        qs.values("sales__username", "sales__first_name")
          .annotate(jumlah=Count("id")).order_by("-jumlah")
    )
    ctx = {
        "total": qs.count(),
        "rekap_status": rekap_status,
        "rekap_sales": rekap_sales,
        "daftar_kunjungan": qs[:100],
        "semua_sales": User.objects.filter(role=User.Role.SALES),
        "status_choices": Kunjungan.Status.choices,
        "f_bulan": bulan, "f_sales": sales_id, "f_status": status,
    }
    return render(request, "report/dashboard_spv.html", ctx)


# --- KUNJUNGAN (Sales) -------------------------------------------------------
@login_required
def kunjungan_list(request):
    if not request.user.is_sales:
        return HttpResponseForbidden("Halaman ini untuk sales.")
    kunjungan = Kunjungan.objects.filter(
        sales=request.user).select_related("instansi")
    return render(request, "report/kunjungan_list.html",
                  {"daftar": kunjungan})


@login_required
def kunjungan_baru(request):
    if not request.user.is_sales:
        return HttpResponseForbidden("Halaman ini untuk sales.")

    if request.method == "POST":
        form = KunjunganForm(request.POST)
        if form.is_valid():
            nama = form.cleaned_data["nama_instansi"].strip()
            instansi, _ = Instansi.objects.get_or_create(
                nama__iexact=nama,
                defaults={
                    "nama": nama,
                    "alamat": form.cleaned_data.get("alamat", ""),
                    "pic": form.cleaned_data.get("pic", ""),
                    "no_wa": form.cleaned_data.get("no_wa", ""),
                    "dibuat_oleh": request.user,
                },
            )
            k = form.save(commit=False)
            k.sales = request.user
            k.instansi = instansi
            k.save()
            FollowUpWA.objects.create(kunjungan=k)
            messages.success(request, "Kunjungan tersimpan.")
            return redirect("kunjungan_list")
    else:
        form = KunjunganForm()

    return render(request, "report/kunjungan_form.html", {
        "form": form,
        "daftar_instansi": Instansi.objects.all(),
    })


@login_required
def kunjungan_edit(request, pk):
    k = get_object_or_404(Kunjungan, pk=pk)
    if k.sales != request.user:
        return HttpResponseForbidden("Bukan kunjungan Anda.")
    if request.method == "POST":
        form = KunjunganForm(request.POST, instance=k)
        if form.is_valid():
            obj = form.save(commit=False)
            obj.sales = request.user
            obj.save()
            messages.success(request, "Kunjungan diperbarui.")
            return redirect("kunjungan_list")
    else:
        form = KunjunganForm(instance=k,
                             initial={"nama_instansi": k.instansi.nama})
    return render(request, "report/kunjungan_form.html", {
        "form": form, "edit": True,
        "daftar_instansi": Instansi.objects.all(),
    })


# --- API cek instansi (notifikasi 'sudah dikunjungi sales lain') -------------
@login_required
def cek_instansi(request):
    nama = request.GET.get("nama", "").strip()
    if not nama:
        return JsonResponse({"ada": False})
    inst = Instansi.objects.filter(nama__iexact=nama).first()
    if not inst:
        return JsonResponse({"ada": False})
    return JsonResponse({
        "ada": True,
        "alamat": inst.alamat, "pic": inst.pic, "no_wa": inst.no_wa,
        "dikunjungi": [
            {"sales": d["sales"], "tanggal": d["tanggal"].strftime("%d/%m/%Y")}
            for d in inst.sudah_dikunjungi_oleh
        ],
    })


# --- BLAST WA ----------------------------------------------------------------
@login_required
def blast_wa(request):
    if not request.user.is_sales:
        return HttpResponseForbidden("Halaman ini untuk sales.")

    kunjungan = Kunjungan.objects.filter(
        sales=request.user
    ).select_related("instansi", "followup").order_by("-tanggal")

    hari_ini = timezone.localdate()
    items = []
    for k in kunjungan:
        links = wa_templates.semua_link_wa(k)
        # Reminder badge: muncul tepat di H+2 / H+5 dari tanggal kunjungan
        selisih = (hari_ini - k.tanggal).days
        reminder = None
        if selisih == 2:
            reminder = "H2"
        elif selisih == 5:
            reminder = "H5"
        items.append({"kunjungan": k, "links": links, "reminder": reminder})

    return render(request, "report/blast_wa.html", {"items": items})


@login_required
def klik_wa(request, pk, tahap):
    """Catat klik link WA (POST/CSRF) -> auto update status, redirect ke WA."""
    if request.method != "POST":
        return redirect("blast_wa")
    k = get_object_or_404(Kunjungan, pk=pk)
    if k.sales != request.user:
        return HttpResponseForbidden("Bukan kunjungan Anda.")
    fu, _ = FollowUpWA.objects.get_or_create(kunjungan=k)
    if tahap in ("H0", "H2", "H5"):
        fu.catat_klik(tahap)
    links = wa_templates.semua_link_wa(k)
    target = links.get(tahap, {}).get("link", "")
    if target:
        return redirect(target)
    return redirect("blast_wa")




# --- DETAIL KUNJUNGAN (drill-down SPV & pemilik) — AI-3 ----------------------
@login_required
def kunjungan_detail(request, pk):
    k = get_object_or_404(
        Kunjungan.objects.select_related("sales", "instansi"), pk=pk)
    boleh = (k.sales == request.user or request.user.is_admin_spv
             or request.user.is_superuser)
    if not boleh:
        return HttpResponseForbidden("Tidak boleh mengakses kunjungan ini.")
    fu = getattr(k, "followup", None)
    return render(request, "report/kunjungan_detail.html",
                  {"k": k, "fu": fu})


# --- EXPORT EXCEL (SPV) ------------------------------------------------------
@login_required
def export_excel(request):
    if not (request.user.is_admin_spv or request.user.is_superuser):
        return HttpResponseForbidden("Hanya SPV/Admin.")

    import openpyxl
    from openpyxl.styles import Font, PatternFill

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Laporan Kunjungan"

    header = ["Tanggal", "Sales", "Instansi", "PIC", "No. WA",
              "Alamat", "Status", "Catatan"]
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="8B1A2B")

    qs = Kunjungan.objects.select_related("sales", "instansi").order_by("-tanggal")
    for k in qs:
        ws.append([
            k.tanggal.strftime("%d/%m/%Y"),
            k.sales.get_full_name() or k.sales.username,
            k.instansi.nama, k.pic or k.instansi.pic,
            k.no_wa or k.instansi.no_wa, k.alamat,
            k.get_status_display(), k.catatan,
        ])

    for col in ws.columns:
        width = max((len(str(c.value)) for c in col if c.value), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(width + 2, 50)

    resp = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    tgl = timezone.localdate().strftime("%Y%m%d")
    resp["Content-Disposition"] = f'attachment; filename="laporan_kunjungan_{tgl}.xlsx"'
    wb.save(resp)
    return resp


# --- AI-7: data notifikasi (dipakai context_processor) ----------------------
def hitung_reminder(user):
    """Jumlah follow-up yang 'waktunya' hari ini (H+2/H+5) untuk sales ini."""
    if not (user.is_authenticated and getattr(user, "is_sales", False)):
        return 0
    hari_ini = timezone.localdate()
    n = 0
    for k in Kunjungan.objects.filter(sales=user).only("tanggal"):
        selisih = (hari_ini - k.tanggal).days
        if selisih in (2, 5):
            n += 1
    return n
